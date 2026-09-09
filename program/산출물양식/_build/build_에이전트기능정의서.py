# -*- coding: utf-8 -*-
"""에이전트(Team) 기능 정의서를 만든다.

근거 — 셋 다 실측·문서에서 가져오고 지어내지 않는다.
  · program/briefing/A-COP_Team모듈_제작_브리핑.html  팀별 manifest·검증 항목
  · program/plan/A-COP_Team모듈_계약사전과_빠진두단계.md  execute() 7단계·계약 필드
  · program/plan/A-COP_구현계획서_v10.md §5·§6  Team 목록·승계 규칙

★담당은 2026-09-09 회의록에서 가져왔다. **Activity 와 Dining 은 담당이 둘씩이라
  그대로 적고 확인 대상으로 표시한다.** 임의로 한 명을 고르지 않는다.

실행: python program/산출물양식/_build/build_에이전트기능정의서.py
"""
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = os.path.dirname(os.path.abspath(__file__))
FORMS = os.path.dirname(HERE)
OUT = os.path.join(FORMS, "[기획] 에이전트 기능 정의서_여행.xlsx")

# ── 서식 ──────────────────────────────────────────────────────────────
NAVY = "1F3864"
HEAD = PatternFill("solid", fgColor=NAVY)
SUB = PatternFill("solid", fgColor="D9E2F3")
MVP = PatternFill("solid", fgColor="FCE4E4")
CUT = PatternFill("solid", fgColor="F2F2F2")
STUB = PatternFill("solid", fgColor="EDEDED")
WARN = PatternFill("solid", fgColor="FFF2CC")
HEADF = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
BOLD = Font(name="맑은 고딕", size=10, bold=True)
BASE = Font(name="맑은 고딕", size=10)
SMALL = Font(name="맑은 고딕", size=9, color="595959")
TITLE = Font(name="맑은 고딕", size=14, bold=True, color=NAVY)
THIN = Side(style="thin", color="B4C6E7")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP = Alignment(vertical="top", wrap_text=True)
CTR = Alignment(vertical="center", horizontal="center", wrap_text=True)


def head(ws, row, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill, cell.font, cell.alignment, cell.border = HEAD, HEADF, CTR, BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30


def put(ws, row, values, fill=None, bold_first=False):
    for i, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=i, value=v)
        cell.font = BOLD if (bold_first and i == 1) else BASE
        cell.alignment = TOP
        cell.border = BOX
        if fill:
            cell.fill = fill


def title(ws, text, sub=None, span=6):
    ws.cell(row=1, column=1, value=text).font = TITLE
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    if sub:
        c = ws.cell(row=2, column=1, value=sub)
        c.font = SMALL
        c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)
        ws.row_dimensions[2].height = 28


wb = openpyxl.Workbook()

# ══════════════════════════════════════════════════════════ 1. 개요
ws = wb.active
ws.title = "개요"
title(ws, "에이전트(Team) 기능 정의서 — 여행 지속관리 CS 플랫폼",
      "A-COP · 2026-09-09 · 근거: 구현계획서 v10 §5·§6, Team 모듈 제작 브리핑, "
      "Team 계약사전. 담당은 2026-09-09 회의록.", span=4)

r = 4
put(ws, r, ["항목", "내용"], bold_first=True)
for c in ws[r]:
    c.fill, c.font = SUB, BOLD
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 96
r += 1
for k, v in [
    ("이 문서가 정하는 것",
     "에이전트(Team) 하나가 무엇을 할 수 있고, 무엇을 입력받아, 어떻게 판정하고, "
     "무엇을 돌려주는가. capability 단위로 적는다."),
    ("이 문서가 안 정하는 것",
     "코어(Core)의 기능. Team 은 side effect 를 실행하지 않으므로 실행·통지·승인은 "
     "코어 몫이다 — 「코어 경계」 시트 참조."),
    ("에이전트 = Team",
     "이 프로젝트에서 에이전트는 Registry 에 등록되는 Team 모듈을 뜻한다. "
     "선언(manifest)·주입(도구)·단일 진입점(execute)·근거(evidence) 넷으로 이루어진다."),
    ("MVP 범위",
     "Activity · Booking Handoff 둘이 필수. Dining · Mobility 는 절삭 가능. "
     "Lodging · Flight 는 등록만 하고 로직을 만들지 않는다."),
    ("판정과 생성을 나눈다",
     "도메인 규칙 판정은 코드가 한다. LLM 은 깨졌을 때 대안을 만드는 데만 쓰고, "
     "만든 대안은 판정을 다시 통과해야 나간다(execute 7단계 ⑤→⑥)."),
    ("근거 의무",
     "모든 판정에 Evidence(출처·확인 시각)를 붙인다. 근거 없는 문장은 계약 단계에서 막힌다."),
    ("상태 표기",
     "[실측] 코드·문서에서 확인 · [설계] 문서에 적혔으나 미구현 · [미정] 안 정해짐"),
]:
    put(ws, r, [k, v], bold_first=True)
    ws.row_dimensions[r].height = 32
    r += 1

r += 1
c = ws.cell(row=r, column=1, value="★ 이 문서를 읽기 전에")
c.font, c.fill = BOLD, WARN
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
r += 1
put(ws, r, ["", "Team 모듈 제작 브리핑(program/briefing/A-COP_Team모듈_제작_브리핑.html)을 "
                "먼저 본다. 이 정의서는 그 브리핑의 팀별 명세를 기능 단위 표로 편 것이다."])
ws.row_dimensions[r].height = 30

# ══════════════════════════════════════════════════════ 2. 에이전트 목록
ws = wb.create_sheet("에이전트 목록")
title(ws, "에이전트(Team) 목록 — manifest 요약",
      "manifest 는 코드 밖에서 읽는 계약이다. Registry 가 이걸 보고 라우팅하고 도구 호출을 막는다.",
      span=9)
COLS = ["Team", "구분", "무엇을 다루나", "team_id", "capabilities",
        "accepted_case_types", "allowed_tools", "knowledge_scope", "담당"]
W = [16, 11, 30, 18, 34, 16, 30, 26, 18]
head(ws, 4, COLS, W)

TEAMS = [
    ("Activity", "MVP 필수",
     "레저 활동 — 골프(티타임)·한강 수상·겨울 스키/눈썰매. 우천 위약금 규정이 문서로 있고 "
     "예약 시각이 고정인 것. 정보 제공 + 예약 판정",
     "activity",
     "activity.check_cancelable\nactivity.request_change\nrefund.calculate",
     "activity",
     "read.booking\nread.policy\nread.place\nread.weather",
     "activity, cancellation,\nrefund, weather",
     "팀모듈1 · 팀모듈3\n★둘 다 액티비티로 적혀 있다", MVP),
    ("Booking Handoff", "MVP 필수",
     "업체 예약 인계 — 기본 동작은 「인계」다. 무엇을 어떻게 바꿔야 하는지 정리해 넘긴다. "
     "실제 변경은 시연 모드(Mock 공급자) 한정",
     "booking",
     "booking.verify\nbooking.prepare_change\nbooking.prepare_cancel\nbooking.status",
     "booking",
     "read.booking\nread.policy\nread.supplier",
     "booking, cancellation,\npenalty, supplier",
     "미정", MVP),
    ("Dining", "절삭 가능",
     "식당 — 그 일정 시각에 여는지와 동행 조건",
     "dining",
     "dining.check_open\ndining.check_conditions\ndining.request_reservation",
     "dining",
     "read.place\nread.policy\nread.booking",
     "dining, opening_hours,\ndietary",
     "팀모듈4 · 검증\n★둘 다 요식으로 적혀 있다", CUT),
    ("Mobility", "절삭 가능",
     "구간 이동 — 환승·막차·여유 시간",
     "mobility",
     "mobility.route\nmobility.status\nmobility.exception",
     "mobility",
     "read.route\nread.transit\nread.policy",
     "mobility, transit,\nroute_exception",
     "팀모듈2", CUT),
    ("Lodging", "등록만",
     "숙박 — 잠긴 예약(locked: true)으로만 취급. 로직 없음",
     "lodging", "—", "—", "—", "—", "—", STUB),
    ("Flight", "등록만",
     "항공 — 잠긴 예약(locked: true)으로만 취급. 로직 없음",
     "flight", "—", "—", "—", "—", "—", STUB),
]
r = 5
for t in TEAMS:
    put(ws, r, list(t[:9]), fill=t[9], bold_first=True)
    ws.row_dimensions[r].height = 74
    r += 1

r += 1
for note in [
    "★ max_steps 는 Team 이 직접 넘겨야 상한이 걸린다 — tools.call(..., budget=self.manifest.max_steps). "
    "안 넘기면 도구 호출 상한이 없는 것과 같다.",
    "★ default_capability 를 안 적으면 Registry 가 고를 수 없다. Activity 는 activity.check_cancelable 로 정해져 있다.",
    "★ accepted_case_types 를 비우면 Case 를 직접 안 받는 팀이 된다.",
]:
    c = ws.cell(row=r, column=1, value=note)
    c.font, c.alignment = SMALL, TOP
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    ws.row_dimensions[r].height = 26
    r += 1

# ══════════════════════════════════════════════════════ 3. 기능 정의 (본표)
ws = wb.create_sheet("기능 정의")
title(ws, "기능 정의 — capability 단위",
      "capability 하나가 기능 하나다. Registry 는 task 의 capability 가 manifest 에 있는지 보고 라우팅한다.",
      span=12)
COLS = ["기능 ID", "에이전트", "capability", "기능명", "하는 일",
        "입력(무엇이 있어야 하나)", "판정 방식", "출력 NextAction",
        "쓰는 도구", "근거(Evidence)", "MVP", "상태"]
W = [11, 15, 24, 18, 40, 28, 20, 20, 22, 24, 8, 9]
head(ws, 4, COLS, W)

FUNCS = [
    # ---------------------------------------------------------- Activity
    ("FN-ACT-01", "Activity", "activity.check_cancelable", "취소 가능 판정",
     "예약 시각과 현재 시각의 차이를 규정 시간과 대조해 지금 취소할 수 있는지 판정한다. "
     "★기본 capability(default_capability)다.",
     "예약 원문(예약 시각·인원) · 취소 규정 · 현재 시각",
     "코드 — LLM 안 부름",
     "RESPOND (가능)\nWAIT_FOR_APPROVAL (위약금 발생)",
     "read.booking\nread.policy",
     "규정 조항 · 예약 원문 · 확인 시각", "●", "[설계]"),
    ("FN-ACT-02", "Activity", "activity.request_change", "변경 접수",
     "날짜 이동·같은 시간대 대체로 바꾸는 요청을 접수한다. 운영일·운영시간·정원 안인지 먼저 본다.",
     "변경 희망 일시 · 인원 · 운영 정보",
     "코드 판정 → LLM 대안 생성 → 코드 재검증",
     "RESPOND\nWAIT_FOR_APPROVAL\nESCALATE (대안 없음)",
     "read.booking\nread.place\nread.policy",
     "운영 공지 · 규정 조항 · 확인 시각", "●", "[설계]"),
    ("FN-ACT-03", "Activity", "refund.calculate", "환급액 계산",
     "남은 시간 구간별 위약금율 표로 환급액을 계산한다. "
     "★규정을 모르면 금액을 만들지 않는다 — 미확인으로 반환한다.",
     "예약 금액 · 취소 시각 · 위약금율 표",
     "코드 — LLM 안 부름",
     "RESPOND\nESCALATE (규정 미확인)",
     "read.policy\nread.booking",
     "위약금율 표 조항 · 계산 근거", "●", "[설계]"),
    ("FN-ACT-W1", "Activity", "(감시)", "날씨 조건 감시",
     "기상청 초단기예보를 재조회해 규정에 기준 시각이 있는 경우만 판정한다"
     "(예: 티타임 2시간 전 강수량). ★코어 스케줄러가 Case 를 연다.",
     "예보 · 규정의 기상 판단 기준 시각",
     "코드 — 규정에 기준이 있을 때만",
     "(Case 생성 → 위 기능으로)",
     "read.weather", "예보 발표 시각 · 관측 시각", "●", "[설계]"),
    # ---------------------------------------------------- Booking Handoff
    ("FN-BKG-01", "Booking Handoff", "booking.verify", "예약 확인",
     "예약이 실재하는지, 우리가 가진 내용과 공급자 기록이 같은지 대조한다.",
     "예약 번호 · 공급자 식별자",
     "코드 — LLM 안 부름",
     "RESPOND\nESCALATE (불일치)",
     "read.booking\nread.supplier",
     "공급자 응답 · 조회 시각", "●", "[설계]"),
    ("FN-BKG-02", "Booking Handoff", "booking.prepare_change", "변경 인계 준비",
     "★기본 동작은 인계다. 무엇을 어떻게 바꿔야 하는지 정리해 "
     "변경 링크 + 차액 + 대안으로 넘긴다. Team 이 실행하지 않는다.",
     "현재 예약 · 희망 변경 · 공급자 규정 · 위임 범위",
     "코드 판정 → LLM 문안 생성 → 코드 재검증",
     "RESPOND (인계)\nWAIT_FOR_APPROVAL (ActionProposal)",
     "read.booking\nread.policy\nread.supplier",
     "공급자 규정 · 차액 산출 근거", "●", "[설계]"),
    ("FN-BKG-03", "Booking Handoff", "booking.prepare_cancel", "취소 인계 준비",
     "취소 조건·위약금을 정리해 인계한다. 실행하지 않는다.",
     "현재 예약 · 취소 규정 · 위임 범위",
     "코드 — LLM 안 부름",
     "RESPOND (인계)\nWAIT_FOR_APPROVAL",
     "read.booking\nread.policy",
     "취소 규정 조항 · 위약금 근거", "●", "[설계]"),
    ("FN-BKG-04", "Booking Handoff", "booking.status", "예약 상태 조회",
     "공급자 쪽 현재 상태를 조회해 돌려준다.",
     "예약 번호",
     "코드 — LLM 안 부름",
     "RESPOND",
     "read.supplier\nread.booking",
     "공급자 응답 · 조회 시각", "●", "[설계]"),
    ("FN-BKG-D1", "Booking Handoff", "(시연 모드 한정)", "위임 범위 대조 후 Mock 실행 제안",
     "시연 모드에서만 위임 범위(금액·종류·되돌림 조건·횟수)를 대조하고 "
     "Mock 공급자 실행을 ActionProposal 로 제안한다. "
     "★idempotency_key 를 반드시 붙인다 — 안 붙이면 같은 변경이 두 번 나갈 수 있다.",
     "위임 범위 선언 · 변경안 · Mock 공급자",
     "코드 — 범위 밖이면 인계로 내린다",
     "WAIT_FOR_APPROVAL (ActionProposal)",
     "read.supplier\nread.policy",
     "위임 범위 조항 · 대조 결과", "●", "[설계]"),
    # ------------------------------------------------------------ Dining
    ("FN-DIN-01", "Dining", "dining.check_open", "영업 여부 판정",
     "★오늘 여는지가 아니라 그 일정 시각에 여는지를 본다. "
     "제공자 영업시간은 오늘부터 7일의 예정이지 조회 순간의 개점 확인이 아니다 — "
     "confirmed_at 과 출처를 함께 남긴다.",
     "장소 식별자 · 일정 시각 · 영업시간 데이터",
     "코드 — 확인 수준을 같이 저장",
     "RESPOND\nRESPOND(미확인 표시)",
     "read.place",
     "출처 · 확인 시각 · 확인 수준", "○", "[설계]"),
    ("FN-DIN-02", "Dining", "dining.check_conditions", "동행 조건 판정",
     "아이 동반 · 할랄 · 채식 같은 조건을 만족하는지 본다.",
     "동행 조건 · 장소 속성",
     "코드 — LLM 안 부름",
     "RESPOND\nRESPOND(미확인 표시)",
     "read.place\nread.policy",
     "장소 속성 출처 · 확인 시각", "○", "[설계]"),
    ("FN-DIN-03", "Dining", "dining.request_reservation", "예약 필요·성립 판정",
     "예약이 필요한 곳인지, 이미 잡혀 있는지 판정한다.",
     "장소 예약 정책 · 기존 예약",
     "코드 판정 → LLM 대안 생성 → 코드 재검증",
     "RESPOND\nWAIT_FOR_INPUT\nESCALATE",
     "read.place\nread.booking",
     "예약 정책 · 예약 원문", "○", "[설계]"),
    # ---------------------------------------------------------- Mobility
    ("FN-MOB-01", "Mobility", "mobility.route", "구간 이동 성립 판정",
     "구간 이동 시간이 일정 간격 안에 들어가는지(여유 시간), 환승이 성립하는지 본다.",
     "출발·도착 · 일정 간격 · 시간표",
     "코드 — 저장한 시간표로 판정",
     "RESPOND\nESCALATE (불성립)",
     "read.route\nread.transit",
     "시간표 출처 · 확인 시각", "○", "[설계]"),
    ("FN-MOB-02", "Mobility", "mobility.status", "운행 상태 조회",
     "해당 구간의 현재 운행 정보를 조회해 돌려준다.",
     "구간 · 시각",
     "코드 — LLM 안 부름",
     "RESPOND",
     "read.transit",
     "운행 정보 출처 · 조회 시각", "○", "[설계]"),
    ("FN-MOB-03", "Mobility", "mobility.exception", "사건 대안 생성",
     "막차를 놓쳤을 때처럼 이동이 깨진 경우 경로·순서를 재배열한다. "
     "★문의와 사건을 가르는 문구 판정이 함정이다 — "
     "「막차가 몇 시인가요」(문의) 대 「막차를 놓쳤어요」(사건). "
     "요청을 분명히 밝히는 문구만 잡고, 신호가 없으면 기본 동작에 맡긴다.",
     "현재 위치 · 남은 일정 · 막차 시각",
     "코드 판정 → 규칙 열거로 대안 → 코드 재검증",
     "RESPOND\nESCALATE (대안 없음)",
     "read.route\nread.transit\nread.policy",
     "시간표 · 규정 · 확인 시각", "○", "[설계]"),
    # ------------------------------------------------------ Lodging/Flight
    ("FN-LDG-00", "Lodging", "(없음)", "등록만",
     "manifest 만 만들고 로직은 없다. 잠긴 예약(locked: true)으로만 취급해 "
     "코어가 다른 일정을 짤 때 고정 항목으로 쓴다.",
     "—", "—", "—", "—", "—", "×", "[설계]"),
    ("FN-FLT-00", "Flight", "(없음)", "등록만",
     "manifest 만 만들고 로직은 없다. 잠긴 예약(locked: true)으로만 취급한다.",
     "—", "—", "—", "—", "—", "×", "[설계]"),
]
r = 5
for f in FUNCS:
    fill = MVP if f[10] == "●" else (STUB if f[10] == "×" else CUT)
    put(ws, r, list(f), fill=fill, bold_first=True)
    ws.row_dimensions[r].height = 76
    r += 1

r += 1
c = ws.cell(row=r, column=1,
            value="MVP 표기 — ● 필수 · ○ 절삭 가능 · × 등록만. "
                  "상태 표기 — [실측] 코드에서 확인 · [설계] 문서에 있으나 미구현 · [미정] 안 정해짐. "
                  "★전 기능이 [설계]다. 2026-09-09 기준 여행용 Team 은 아직 코드가 없다.")
c.font, c.alignment, c.fill = BOLD, TOP, WARN
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
ws.row_dimensions[r].height = 32

# ══════════════════════════════════════════════════════ 4. 처리 흐름
ws = wb.create_sheet("처리 흐름")
title(ws, "execute() 7단계 — 모든 에이전트가 같은 순서로 돈다",
      "브리핑은 다섯 칸으로 적었으나 v10 §5 는 일곱을 요구한다. ⑤ 생성과 ⑥ 재검증이 빠져 있었다.",
      span=5)
head(ws, 4, ["단계", "이름", "누가 하나", "하는 일", "안 지키면"],
     [8, 14, 14, 58, 46])
STEPS = [
    ("1", "가드", "코드",
     "team_id 가 내 것인지 · capability 가 manifest 에 있는지 · context 가 degraded 인지 본다",
     "다른 팀 일을 받아 처리하거나, 자료가 없는 채로 판정한다"),
    ("2", "준비", "코드",
     "seen 집합으로 같은 도구·같은 인자 재호출을 막고, context 의 근거를 누적 시작한다",
     "같은 조회를 반복해 비용이 늘고 도구 상한이 무의미해진다"),
    ("3", "조회", "코드",
     "tools.call(name, context, args, allowed_tools, seen, budget=max_steps) 로 필요한 것만 읽는다",
     "budget 을 안 넘기면 도구 호출 상한이 안 걸린다"),
    ("4", "판정", "코드",
     "도메인 규칙은 전부 여기. LLM 을 부르지 않는다. 성립하면 바로 반환한다",
     "규칙을 LLM 에 맡기면 같은 입력에 다른 답이 나온다"),
    ("⑤", "생성", "LLM",
     "④가 깨졌을 때만 대안을 만든다. llm 이 없으면 건너뛴다",
     "④→⑦ 로 가면 「안 된다」만 말하는 제품이 된다"),
    ("⑥", "재검증", "코드",
     "★④와 같은 함수를 다시 부른다. 통과한 대안만 남긴다",
     "다른 함수를 쓰면 「생성용 느슨한 판정」이 생겨 원칙이 무너진다. "
     "⑤→⑦ 로 가면 검증 안 된 대안이 고객에게 간다"),
    ("7", "반환", "코드",
     "TeamResult 로 돌려준다. 근거(Evidence)와 NextAction 을 함께 싣는다",
     "근거 없는 답변은 계약 단계에서 막힌다"),
]
r = 5
for s in STEPS:
    fill = WARN if s[0] in ("⑤", "⑥") else None
    put(ws, r, list(s), fill=fill, bold_first=True)
    ws.row_dimensions[r].height = 40
    r += 1

r += 1
c = ws.cell(row=r, column=1,
            value="★ ⑥이 이 아키텍처의 실체다. v9 에서는 이 원칙을 별도 Team(response_review)으로 뒀는데 "
                  "v10 이 각 Team 안으로 넣었다. 그래서 팀마다 판정 함수를 두 번 부르는 모양이 된다.")
c.font, c.alignment, c.fill = BOLD, TOP, WARN
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
ws.row_dimensions[r].height = 32

# ══════════════════════════════════════════════════════ 5. 코어 경계
ws = wb.create_sheet("코어 경계")
title(ws, "에이전트가 안 하는 것 — 코어 몫",
      "Team 은 side effect 를 실행하지 않는다(v10 §6). 제안까지가 Team 이고 실행은 코어다.",
      span=3)
head(ws, 4, ["일", "누가", "왜"], [34, 16, 74])
BOUND = [
    ("도메인 판정 · 대안 생성 · 근거 수집", "Team", "자기 객체의 규칙은 그 팀이 가장 잘 안다"),
    ("전체 일정 정합성(시간 충돌·예산·이동 여유)", "코어 검증 층",
     "한 팀이 자기 객체만 보면 되도록 나눈 분리다. 팀이 전체를 보면 팀마다 같은 검사를 갖게 된다"),
    ("Case 생성 — 고객 요청 · 감시 사건", "코어", "감시가 Case 를 여는 것이 여행 도메인의 새 구조다"),
    ("실제 실행(예약 변경·취소·결제)", "코어 Action 층",
     "Team 은 ActionProposal 로 제안만 한다. 실행 권한을 팀에 주면 승인 경계가 무너진다"),
    ("고객 통지 발송", "코어", "Team 은 무엇을 보낼지 만들고, 보내는 것은 코어가 한다"),
    ("승인 대기 · 되돌림", "코어", "승인 상태와 일정 버전은 코어가 소유한다"),
    ("사람 인계(escalated)", "코어", "Team 은 escalated 를 반환하고, 넘기는 처리는 코어가 한다"),
]
r = 5
for b in BOUND:
    put(ws, r, list(b), bold_first=True)
    ws.row_dimensions[r].height = 34
    r += 1

# ══════════════════════════════════════════════════════ 6. 미해결
ws = wb.create_sheet("미해결")
title(ws, "정하지 않은 것 — 만들기 전에 정해야 한다",
      "지어내지 않고 미정으로 둔 항목이다.", span=3)
head(ws, 4, ["항목", "무엇이 문제인가", "누가 정하나"], [30, 74, 20])
OPEN = [
    ("Activity 담당이 둘",
     "2026-09-09 회의록에 팀모듈1과 팀모듈3이 모두 「액티비티 모듈」로 적혀 있다. "
     "같은 모듈을 둘이 하는 것인지, 한쪽이 다른 모듈인지 확인이 필요하다.", "팀"),
    ("Dining 담당이 둘",
     "팀모듈4가 요식업 데이터·모듈 제안서를, 검증이 Dining 모듈 구조 설계를 하고 있다. "
     "역할 경계가 문서에 없다.", "팀"),
    ("Booking Handoff 담당 미정",
     "회의록에 담당이 나오지 않는다. MVP 필수 둘 중 하나다.", "팀"),
    ("Activity 의 범위",
     "판정 규칙은 날씨 조건이고 감시 소스는 기상청인데, v10 예시 목록에 공연·쿠킹 클래스 같은 "
     "실내 상품이 섞여 있다. 실내를 넣으면 감시 소스가 둘로 갈려 「객체 종류별로 나눈다」는 "
     "분할 원칙과 충돌한다.", "팀"),
    ("Case 에 타는 도메인 객체 id 규칙",
     "코어는 이미 도메인 중립이다(business_subject 라는 문자열 한 칸). "
     "문제는 그 칸에 무엇을 넣는지 규칙이 없다는 것이다.", "코어"),
    ("A2A · MCP 자리",
     "v10 에 A2A·MCP 문자열이 0회다. 부트캠프 요구사항 대응표도 없다. "
     "기능 판단이 아니라 채점 항목이다.", "팀 · 코어"),
    ("공통 헬퍼 위치",
     "판정 함수를 두 번 부르는 모양이 모든 팀에 생긴다. 공통 헬퍼를 어디에 둘지 정해야 "
     "팀마다 다르게 구현되는 것을 막는다.", "코어"),
]
r = 5
for o in OPEN:
    put(ws, r, list(o), fill=WARN, bold_first=True)
    ws.row_dimensions[r].height = 46
    r += 1

for sheet in wb.worksheets:
    sheet.freeze_panes = "A5"
    sheet.sheet_view.showGridLines = False
wb["개요"].freeze_panes = "A4"

wb.save(OUT)
print(f"{os.path.basename(OUT)}  ·  시트 {len(wb.worksheets)}개  ·  "
      f"기능 {len(FUNCS)}행 · 에이전트 {len(TEAMS)}개 · 미해결 {len(OPEN)}건")
